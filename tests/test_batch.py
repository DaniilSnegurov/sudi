"""Пакет: безопасность чтения, частичный сбой, копии, сохранение типов в xlsx."""

import zipfile

import openpyxl
import pytest

from courtdocs.export import write_all
from courtdocs.models import FAILED
from courtdocs.pipeline import run
from courtdocs.registry import safe_extract

NS = "http://www.fssprus.ru/namespace/order/2017/2"


def _xml(**fields) -> str:
    base = {
        "Id": "1", "DocType": "O_IP_ACT_END_END", "DocDate": "2033-06-10", "DebtorType": "2",
        "DbtrName": "Пробов Назар Назарович", "IdDocNo": "0012/2033", "IdDebtSum": "100.00",
    }
    base.update(fields)
    body = "".join(f"<f:{k}>{v}</f:{k}>" for k, v in base.items() if v is not None)
    return f'<?xml version="1.0" encoding="UTF-8"?><f:OIp xmlns:f="{NS}">{body}</f:OIp>'


def test_req015_broken_file_does_not_stop_batch(tmp_path):
    (tmp_path / "good.xml").write_text(_xml(), encoding="utf-8")
    (tmp_path / "broken.xml").write_text("<f:OIp", encoding="utf-8")
    (tmp_path / "other.xml").write_text('<x:Doc xmlns:x="urn:other"/>', encoding="utf-8")
    batch = run(tmp_path)
    by_name = {d.path: d for d in batch.documents}
    assert by_name["good.xml"].status != FAILED
    assert {i.code for i in by_name["broken.xml"].issues} == {"XML_PARSE_ERROR"}
    assert {i.code for i in by_name["other.xml"].issues} == {"UNSUPPORTED_XML_SCHEMA"}
    # неудачные файлы остаются в таблице с путём и попадают в отчёт
    write_all(batch, tmp_path / "out")
    ws = openpyxl.load_workbook(tmp_path / "out" / "fssp.xlsx").active
    assert [r[0] for r in ws.iter_rows(min_row=2, values_only=True)] == ["broken.xml", "good.xml", "other.xml"]


def test_external_entities_are_rejected(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("секрет", encoding="utf-8")
    payload = (
        f'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "{secret.as_uri()}">]>'
        f'<f:OIp xmlns:f="{NS}"><f:DbtrName>&x;</f:DbtrName></f:OIp>'
    )
    (tmp_path / "xxe.xml").write_text(payload, encoding="utf-8")
    doc = run(tmp_path, include=()).documents
    xxe = next(d for d in doc if d.path == "xxe.xml")
    assert xxe.status == FAILED and "секрет" not in str(xxe.fields)


def test_missing_core_field_and_unknown_type(tmp_path):
    (tmp_path / "a.xml").write_text(_xml(DocType="O_IP_NEW_KIND", IdDebtSum=None), encoding="utf-8")
    doc = run(tmp_path).documents[0]
    assert {i.code for i in doc.issues} == {"UNKNOWN_DOC_TYPE", "MISSING_CORE_FIELD"}
    assert doc.value("DocType") == "O_IP_NEW_KIND" and doc.value("DocType2") is None


def test_exact_duplicates_are_processed_once(tmp_path):
    (tmp_path / "a.xml").write_text(_xml(), encoding="utf-8")
    (tmp_path / "copy").mkdir()
    (tmp_path / "copy" / "b.xml").write_text(_xml(), encoding="utf-8")
    batch = run(tmp_path)
    assert [d.path for d in batch.documents] == ["a.xml"]
    assert batch.documents[0].duplicates == ["copy/b.xml"]


def test_req017_leading_zero_survives_xlsx_roundtrip(tmp_path):
    (tmp_path / "a.xml").write_text(_xml(IdDeloNo="0012/2033", IdType="04"), encoding="utf-8")
    write_all(run(tmp_path), tmp_path / "out")
    ws = openpyxl.load_workbook(tmp_path / "out" / "fssp.xlsx").active
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["IdDocNo"] == "0012/2033" and row["IdDeloNo"] == "0012/2033" and row["IdType"] == "04"
    assert row["IdDebtSum"] == 100.0 and row["DebtorType"] == 2
    assert row["rub_deb"] is None  # отсутствие значения — пустая ячейка, не ноль


def test_zip_path_traversal_is_rejected(tmp_path):
    archive = tmp_path / "evil.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("../outside.xml", "<a/>")
    with pytest.raises(ValueError):
        safe_extract(archive, tmp_path / "unpacked")
    assert not (tmp_path / "outside.xml").exists()


def test_zip_package_is_processed(tmp_path):
    archive = tmp_path / "pack.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("fssp/a.xml", _xml())
    batch = run(archive, tmp_path / "work")
    assert [d.path for d in batch.documents] == ["fssp/a.xml"]
