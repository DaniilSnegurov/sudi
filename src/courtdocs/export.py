"""Экспорт: таблица ФССП, отчёт о проблемах, машиночитаемый результат."""

from __future__ import annotations

import csv
import json
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .judicial import COLUMNS as JUDICIAL_COLUMNS
from .models import DOC_NEEDS_REVIEW, FAILED, PROCESSED, to_jsonable
from .pipeline import BatchResult

FSSP_COLUMNS = (
    "FileName", "DocType", "DocType2", "DebtorType", "DocDate", "IdDebtText", "IdDocNo", "IdDocDate",
    "IdDeloNo", "IdDeloDate", "IdDebtSum", "IpNo", "IdType", "DbtrAdr", "street", "dom", "kv", "DbtrName",
    "date_start", "date_end", "rub_deb", "rub_peni", "rub_poshlina", "rub_post",
)
# В образце xlsx коды типов записаны числами, остальные реквизиты — текстом
_INT_CODES = ("DebtorType", "IdType")

ISSUE_COLUMNS = (
    "Документ", "Файл", "Статус", "Этап", "Поле", "Код", "Причина", "Источник", "Кандидаты",
    "Влияние на экспорт", "Рекомендуемое действие", "Важность",
)
_STATUS_RU = {PROCESSED: "без выявленных проблем", DOC_NEEDS_REVIEW: "требует проверки", FAILED: "не обработан"}


def _write_cell(ws, row: int, col: int, name: str, value) -> None:
    cell = ws.cell(row=row, column=col)
    if value is None:
        cell.number_format = "@"
        return
    if isinstance(value, Decimal):
        cell.value, cell.number_format = float(value), "#,##0.00"
    elif isinstance(value, date):
        cell.value, cell.number_format = value, "yyyy-mm-dd"
    elif isinstance(value, int):
        cell.value, cell.number_format = value, "0"
    elif name in _INT_CODES and str(value).strip().isdigit() and not str(value).strip().startswith("0"):
        cell.value = int(str(value).strip())
    else:
        cell.value, cell.number_format = ILLEGAL_CHARACTERS_RE.sub("", str(value)), "@"
        cell.data_type = "s"  # ведущие нули и номера не превращаются в числа


def _csv_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return f"{value:.2f}"
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _rows(batch: BatchResult, profile: str, columns: tuple[str, ...], path_column: str) -> list[dict]:
    # у необработанного документа в строке остаётся путь файла
    return [{c: d.path if c == path_column else d.value(c) for c in columns} for d in batch.by_profile(profile)]


def fssp_rows(batch: BatchResult) -> list[dict]:
    return _rows(batch, "fssp", FSSP_COLUMNS, "FileName")


def judicial_rows(batch: BatchResult) -> list[dict]:
    return _rows(batch, "judicial", JUDICIAL_COLUMNS, "Файл")


def _write_table(rows: list[dict], columns: tuple[str, ...], out_dir: Path, name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for c, column in enumerate(columns, 1):
        ws.cell(row=1, column=c, value=column)
    for r, row in enumerate(rows, 2):
        for c, column in enumerate(columns, 1):
            _write_cell(ws, r, c, column, row[column])
    wb.save(out_dir / f"{name}.xlsx")

    with open(out_dir / f"{name}.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        for row in rows:
            w.writerow([_csv_value(row[c]) for c in columns])


def write_fssp(batch: BatchResult, out_dir: Path) -> None:
    _write_table(fssp_rows(batch), FSSP_COLUMNS, out_dir, "fssp")


def write_judicial(batch: BatchResult, out_dir: Path) -> None:
    _write_table(judicial_rows(batch), JUDICIAL_COLUMNS, out_dir, "judicial")


def write_issues(batch: BatchResult, out_dir: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Проблемы"
    ws.append(ISSUE_COLUMNS)
    for d in batch.documents:
        for i in d.issues:
            ws.append([
                d.doc_id, d.path, _STATUS_RU[d.status], i.stage, i.field or "", i.code, i.message,
                i.location or "", "; ".join(i.candidates), i.export_effect, i.action,
                "проверить" if i.severity == "review" else "к сведению",
            ])

    ws2 = wb.create_sheet("Сводка")
    docs = batch.documents
    ws2.append(["Показатель", "Количество"])
    for label, n in (
        ("Всего файлов в пакете", len(batch.entries)),
        ("Документов (строк результата)", len(docs)),
        ("Без выявленных проблем", sum(d.status == PROCESSED for d in docs)),
        ("Требуют проверки", sum(d.status == DOC_NEEDS_REVIEW for d in docs)),
        ("Не обработаны", sum(d.status == FAILED for d in docs)),
        ("Парные PDF", sum(e.paired_with is not None for e in batch.entries)),
        ("Точные копии файлов", sum(e.duplicate_of is not None for e in batch.entries)),
        ("Прочие файлы (не документы)", sum(e.kind == "other" and not e.duplicate_of for e in batch.entries)),
    ):
        ws2.append([label, n])

    ws3 = wb.create_sheet("Файлы")
    ws3.append(["Файл", "Роль", "Статус", "SHA-256"])
    status = {d.path: _STATUS_RU[d.status] for d in docs}
    for e in batch.entries:
        if e.duplicate_of:
            role = f"копия {e.duplicate_of}"
        elif e.paired_with:
            role = f"парный PDF к {e.paired_with}"
        elif e.kind == "other":
            role = "не документ, пропущен"
        else:
            role = "основной документ"
        ws3.append([e.rel_path, role, status.get(e.rel_path, ""), e.sha256])
    wb.save(out_dir / "issues.xlsx")


def write_json(batch: BatchResult, out_dir: Path) -> None:
    data = {"root": str(batch.root), "documents": [to_jsonable(d) for d in batch.documents]}
    (out_dir / "result.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")


def write_all(batch: BatchResult, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    write_fssp(batch, out_dir)
    write_judicial(batch, out_dir)
    write_issues(batch, out_dir)
    write_json(batch, out_dir)
